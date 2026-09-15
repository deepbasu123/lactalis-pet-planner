"""tests/test_excel_source.py

TDD tests for backend.excel_source, the real-data loader for the
"PET Traffic Lights" workbook.

No real/sensitive workbook is used or committed here -- _build_fixture_workbook
constructs a tiny, fully-synthetic .xlsx in the SAME sheet/column shape
(verified against the real 07.09.2026 export) with small, fake numbers, so
these tests are fully offline and portable, matching the rest of the suite.
"""
from __future__ import annotations

import datetime

import openpyxl
import pandas as pd
import pytest

from backend import excel_source as ex

# ---------------------------------------------------------------------------
# Fixture workbook: 3 SKUs, 6-week horizon, same shape as the real file.
# ---------------------------------------------------------------------------

_HORIZON = 6
_START = datetime.date(2026, 9, 7)  # a real Monday, matches the source file's anchor
_MAINT_OFFSET = 2  # 0-indexed -> horizon_index 3

# (row, code, description) -- mirrors the real "PET " sheet's rows 15+.
# Row 16 (between two active SKUs, not after them) is inactive -- this
# specifically exercises the "does a skipped row shift everything after it
# onto the wrong worksheet row" bug class, not just "is a trailing row
# excluded".
_FIXTURE_SKUS = [
    (15, 60444, "PAULS ZYMIL FLAV MILK CHOC 400ML**"),   # has real MLOR row
    (17, 228500, "PAULS PLUS  CHOCOLATE Flavoured Milk  400ML"),  # MLOR missing -> fallback
    (18, 61108, "OAK UHT CHOCOLATE 500ML"),               # has real MLOR row, 500ml
]

_MEASURES = [
    "Forecast", "Sales Order", "DistrDemand (Planned)",
    "DistrDemand (TLB Confirmed)", "Stock on Hand (Projected)",
]

# An inactive row sitting BETWEEN two active SKUs (16 sits between 15 and 17)
# -- must be excluded entirely (not counted, not defaulted to "Active"), and
# must NOT cause 228500/61108's plan quantities to be read off the wrong row.
_INACTIVE_ROW = 16
_INACTIVE_SKU = 999999

# Deliberately fabricated shelf-life/MLOR pairs, chosen to be nothing like the
# real source file's real per-SKU figures (this is a fixture with fake SKU
# codes reused for realism -- only the loader's *logic* is under test here).
_FAKE_MLOR = {60444: (112, 28), 61108: (140, 35)}

# code -> {week offset -> planned qty}; every other cell is left blank (0.0).
# Values are deliberately non-round/arbitrary so they cannot be mistaken for
# real production batch sizes copied out of the source workbook.
_PLAN_CELLS: dict[int, dict[int, float]] = {
    60444: {0: 37_219.0, 3: 84_105.0},
    228500: {1: 52_988.0},
    61108: {0: 19_340.0, 5: 61_177.0},
}


def _build_fixture_workbook(path) -> None:
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    dates = [_START + datetime.timedelta(weeks=i) for i in range(_HORIZON)]

    # ---- "PET " sheet ----
    pet = wb.create_sheet("PET ")
    for i, d in enumerate(dates):
        col = ex._PET_WEEK_FIRST_COL + i
        pet.cell(row=ex._PET_WEEK_DATE_ROW, column=col,
                  value=datetime.datetime(d.year, d.month, d.day))
    pet.cell(row=ex._PET_MAINT_ROW, column=ex._PET_WEEK_FIRST_COL + _MAINT_OFFSET,
              value=ex._PET_MAINT_MARKER)

    for row, code, desc in _FIXTURE_SKUS:
        pet.cell(row=row, column=ex._PET_STATUS_COL, value="Active")
        pet.cell(row=row, column=ex._PET_CODE_COL, value=code)
        pet.cell(row=row, column=ex._PET_DESC_COL, value=desc)
        for offset, qty in _PLAN_CELLS.get(code, {}).items():
            pet.cell(row=row, column=ex._PET_WEEK_FIRST_COL + offset, value=qty)

    # An inactive row in the candidate range -- must be skipped entirely.
    pet.cell(row=_INACTIVE_ROW, column=ex._PET_STATUS_COL, value="Discontinued")
    pet.cell(row=_INACTIVE_ROW, column=ex._PET_CODE_COL, value=_INACTIVE_SKU)
    pet.cell(row=_INACTIVE_ROW, column=ex._PET_DESC_COL, value="RETIRED PRODUCT 500ML")

    # ---- "MLOR" sheet (228500 deliberately omitted -> exercises fallback) ----
    mlor = wb.create_sheet("MLOR")
    mlor.cell(row=3, column=1, value="Product")  # header row, positions only matter from row 4
    for i, (code, (shelf, mlor_days)) in enumerate(_FAKE_MLOR.items()):
        r = 4 + i
        mlor.cell(row=r, column=ex._MLOR_CODE_COL, value=code)
        mlor.cell(row=r, column=ex._MLOR_SHELF_COL, value=shelf)
        mlor.cell(row=r, column=ex._MLOR_MLOR_COL, value=mlor_days)

    # ---- "SNP Dump" sheet ----
    snp = wb.create_sheet("SNP Dump")
    for i, d in enumerate(dates):
        snp.cell(row=1, column=ex._SNP_WEEK_FIRST_COL + i, value=d.strftime("%d.%m.%Y"))
    r = 1
    for code in (60444, 228500, 61108):
        for measure in _MEASURES:
            for loc in ("Total", "ROW", "VIC2"):
                r += 1
                snp.cell(row=r, column=ex._SNP_CODE_COL, value=code)
                snp.cell(row=r, column=ex._SNP_MEASURE_COL, value=measure)
                snp.cell(row=r, column=ex._SNP_LOCATION_COL, value=loc)
                if loc != "Total":
                    continue  # loader only reads the 'Total' location row
                snp.cell(row=r, column=ex._SNP_INITIAL_COL, value=1000 + code % 100)
                for i2 in range(_HORIZON):
                    snp.cell(row=r, column=ex._SNP_WEEK_FIRST_COL + i2,
                              value=137 * (i2 + 1) + code % 47)

    wb.save(path)


@pytest.fixture()
def fixture_path(tmp_path):
    path = tmp_path / "fixture_pet_traffic_lights.xlsx"
    _build_fixture_workbook(path)
    return str(path)


@pytest.fixture()
def dataset(fixture_path):
    return ex.load_dataset(fixture_path, horizon_weeks=_HORIZON)


# ---------------------------------------------------------------------------
# Pure helper tests
# ---------------------------------------------------------------------------

class TestPureHelpers:
    def test_infer_pack_size_400(self):
        assert ex.infer_pack_size_ml("PAULS ZYMIL FLAV MILK CHOC 400ML**") == 400

    def test_infer_pack_size_500_mixed_case(self):
        assert ex.infer_pack_size_ml("OAK PLUS FLAVOURED MILK NAS CHOC 6x500ml") == 500

    def test_infer_pack_size_raises_when_absent(self):
        with pytest.raises(ValueError):
            ex.infer_pack_size_ml("SOME PRODUCT WITH NO SIZE")

    def test_iso_week_key_matches_source_anchor(self):
        assert ex.iso_week_key(datetime.date(2026, 9, 7)) == "2026-W37"

    def test_iso_week_key_year_rollover(self):
        assert ex.iso_week_key(datetime.date(2027, 1, 4)) == "2027-W01"

    def test_clean_description_strips_trailing_stars_and_collapses_whitespace(self):
        assert ex._clean_description("PAULS  ZYMIL\xa0CHOC 400ML**") == "PAULS ZYMIL CHOC 400ML"

    def test_clean_description_handles_none(self):
        assert ex._clean_description(None) == ""


# ---------------------------------------------------------------------------
# load_dataset() -- shape and correctness
# ---------------------------------------------------------------------------

class TestLoadDatasetShape:
    def test_returns_all_six_tables(self, dataset):
        assert set(dataset.keys()) == {
            "sku", "week", "parameter", "demand", "plan_line", "opening_stock",
        }

    def test_three_skus(self, dataset):
        assert len(dataset["sku"]) == 3

    def test_six_weeks(self, dataset):
        assert len(dataset["week"]) == _HORIZON

    def test_demand_rows_is_skus_times_weeks(self, dataset):
        assert len(dataset["demand"]) == 3 * _HORIZON

    def test_plan_line_rows_is_skus_times_weeks(self, dataset):
        assert len(dataset["plan_line"]) == 3 * _HORIZON

    def test_opening_stock_one_row_per_sku(self, dataset):
        assert len(dataset["opening_stock"]) == 3

    def test_ten_parameters_reused_from_data_gen(self, dataset):
        assert len(dataset["parameter"]) == 10


class TestSkuTable:
    def test_priority_matches_planner_row_order(self, dataset):
        sku_df = dataset["sku"].set_index("sku_code")
        assert sku_df.loc["60444", "priority"] == 1
        assert sku_df.loc["228500", "priority"] == 2
        assert sku_df.loc["61108", "priority"] == 3

    def test_pack_size_inferred_from_description(self, dataset):
        sku_df = dataset["sku"].set_index("sku_code")
        assert sku_df.loc["60444", "pack_size_ml"] == 400
        assert sku_df.loc["228500", "pack_size_ml"] == 400
        assert sku_df.loc["61108", "pack_size_ml"] == 500

    def test_mlor_used_directly_when_present(self, dataset):
        sku_df = dataset["sku"].set_index("sku_code")
        shelf_60444, mlor_60444 = _FAKE_MLOR[60444]
        shelf_61108, mlor_61108 = _FAKE_MLOR[61108]
        assert sku_df.loc["60444", "shelf_life_days"] == shelf_60444
        assert sku_df.loc["60444", "mlor_days"] == mlor_60444
        assert sku_df.loc["61108", "shelf_life_days"] == shelf_61108
        assert sku_df.loc["61108", "mlor_days"] == mlor_61108

    def test_mlor_imputed_from_fallback_when_missing(self, dataset):
        """228500 has no MLOR row in the fixture -- must inherit 60444's
        values via _MLOR_FALLBACK, not a fabricated placeholder."""
        sku_df = dataset["sku"].set_index("sku_code")
        shelf_60444, mlor_60444 = _FAKE_MLOR[60444]
        assert sku_df.loc["228500", "shelf_life_days"] == shelf_60444
        assert sku_df.loc["228500", "mlor_days"] == mlor_60444

    def test_max_cover_weeks_computed_from_shelf_and_mlor(self, dataset):
        sku_df = dataset["sku"].set_index("sku_code")
        shelf, mlor_days = _FAKE_MLOR[60444]
        assert sku_df.loc["60444", "max_cover_weeks"] == round((shelf - mlor_days) / 7.0, 1)

    def test_inactive_sku_is_excluded_entirely(self, dataset):
        """A row with a non-Active status in the candidate range must not
        appear in the loaded SKU table at all -- not included, and never
        silently defaulted to status='Active'."""
        assert str(_INACTIVE_SKU) not in set(dataset["sku"]["sku_code"])

    def test_status_active(self, dataset):
        assert (dataset["sku"]["status"] == "Active").all()

    def test_description_cleaned(self, dataset):
        sku_df = dataset["sku"].set_index("sku_code")
        assert sku_df.loc["60444", "description"] == "PAULS ZYMIL FLAV MILK CHOC 400ML"


class TestWeekTable:
    def test_week_keys_and_horizon_index_sequential(self, dataset):
        week_df = dataset["week"].sort_values("horizon_index")
        assert list(week_df["horizon_index"]) == list(range(1, _HORIZON + 1))
        assert week_df.iloc[0]["week_key"] == "2026-W37"

    def test_maintenance_flagged_at_source_offset(self, dataset):
        week_df = dataset["week"].sort_values("horizon_index").reset_index(drop=True)
        assert week_df.loc[_MAINT_OFFSET, "maintenance_type"] == "Full"
        others = week_df.drop(index=_MAINT_OFFSET)
        assert (others["maintenance_type"] == "None").all()

    def test_first_three_weeks_locked(self, dataset):
        week_df = dataset["week"].sort_values("horizon_index").reset_index(drop=True)
        assert list(week_df["is_locked"][:3]) == [True, True, True]
        assert not week_df["is_locked"][3:].any()

    def test_week_commencing_is_a_date(self, dataset):
        week_df = dataset["week"]
        assert isinstance(week_df.iloc[0]["week_commencing"], datetime.date)


class TestDemandTable:
    def test_values_match_source_for_first_week(self, dataset):
        demand_df = dataset["demand"]
        row = demand_df[(demand_df["sku_code"] == "60444") & (demand_df["week_key"] == "2026-W37")].iloc[0]
        expected = 137 * 1 + 60444 % 47
        assert row["forecast"] == expected
        assert row["sales_order"] == expected
        assert row["distr_demand_planned"] == expected
        assert row["distr_demand_tlb"] == expected

    def test_values_vary_by_week(self, dataset):
        demand_df = dataset["demand"]
        sub = demand_df[demand_df["sku_code"] == "61108"].sort_values("week_key")
        assert sub["forecast"].nunique() == _HORIZON  # every week's synthetic value is distinct


class TestPlanLineTable:
    def test_populated_cells_match_source(self, dataset):
        plan_df = dataset["plan_line"]
        row = plan_df[(plan_df["sku_code"] == "60444") & (plan_df["week_key"] == "2026-W37")].iloc[0]
        assert row["planned_qty"] == 37_219.0
        assert row["orig_qty"] == 37_219.0

    def test_blank_cells_load_as_zero_not_missing(self, dataset):
        plan_df = dataset["plan_line"]
        row = plan_df[(plan_df["sku_code"] == "60444") & (plan_df["week_key"] == "2026-W38")].iloc[0]
        assert row["planned_qty"] == 0.0
        assert row["orig_qty"] == 0.0

    def test_orig_qty_equals_planned_qty_everywhere(self, dataset):
        plan_df = dataset["plan_line"]
        assert (plan_df["planned_qty"] == plan_df["orig_qty"]).all()

    def test_second_sku_second_week_value(self, dataset):
        plan_df = dataset["plan_line"]
        row = plan_df[(plan_df["sku_code"] == "228500") & (plan_df["week_key"] == "2026-W38")].iloc[0]
        assert row["planned_qty"] == 52_988.0

    def test_no_row_shift_regression_after_skipped_inactive_row(self, dataset):
        """Regression test: an earlier version derived each SKU's worksheet
        row from its priority (_PET_SKU_ROWS[priority - 1]), which is only
        correct if every candidate row is active. With row 16 (between
        60444 and 228500) inactive, priority 2/3 no longer equal row
        position 2/3 -- 228500 (row 17) and 61108 (row 18) must still read
        their OWN cells, not row 16's (blank) or each other's."""
        plan_df = dataset["plan_line"]

        row_228500 = plan_df[
            (plan_df["sku_code"] == "228500") & (plan_df["week_key"] == "2026-W38")
        ].iloc[0]
        assert row_228500["planned_qty"] == 52_988.0  # from real row 17, not row 16 (blank)

        row_61108 = plan_df[
            (plan_df["sku_code"] == "61108") & (plan_df["week_key"] == "2026-W37")
        ].iloc[0]
        assert row_61108["planned_qty"] == 19_340.0  # from real row 18, not row 17's data


class TestOpeningStockTable:
    def test_initial_value_matches_source(self, dataset):
        opening_df = dataset["opening_stock"].set_index("sku_code")
        assert opening_df.loc["60444", "opening_ea"] == 1000 + 60444 % 100
        assert opening_df.loc["228500", "opening_ea"] == 1000 + 228500 % 100
        assert opening_df.loc["61108", "opening_ea"] == 1000 + 61108 % 100


# ---------------------------------------------------------------------------
# Error handling -- fail loudly on unexpected shape, never guess
# ---------------------------------------------------------------------------

class TestErrorHandling:
    def test_missing_sheet_raises_keyerror(self, tmp_path):
        wb = openpyxl.Workbook()
        wb.save(tmp_path / "empty.xlsx")
        with pytest.raises(KeyError):
            ex.load_dataset(str(tmp_path / "empty.xlsx"))

    def test_missing_demand_measure_raises_keyerror(self, tmp_path):
        """A workbook with SKUs but no SNP Dump measure rows must fail loudly,
        not silently zero-fill real demand data."""
        wb = openpyxl.Workbook()
        wb.remove(wb.active)
        pet = wb.create_sheet("PET ")
        pet.cell(row=13, column=ex._PET_WEEK_FIRST_COL,
                  value=datetime.datetime(2026, 9, 7))
        pet.cell(row=15, column=1, value="Active")
        pet.cell(row=15, column=ex._PET_CODE_COL, value=60444)
        pet.cell(row=15, column=ex._PET_DESC_COL, value="TEST PRODUCT 400ML")
        wb.create_sheet("MLOR")
        snp = wb.create_sheet("SNP Dump")  # date header present, but no measure rows
        snp.cell(row=1, column=ex._SNP_WEEK_FIRST_COL, value="07.09.2026")
        path = tmp_path / "incomplete.xlsx"
        wb.save(path)
        with pytest.raises(KeyError):
            ex.load_dataset(str(path), horizon_weeks=1)

    def test_no_active_skus_raises_valueerror(self, tmp_path):
        wb = openpyxl.Workbook()
        wb.remove(wb.active)
        wb.create_sheet("PET ")
        wb.create_sheet("MLOR")
        wb.create_sheet("SNP Dump")
        path = tmp_path / "no_skus.xlsx"
        wb.save(path)
        with pytest.raises(ValueError):
            ex.load_dataset(str(path))

    def test_num_raises_on_non_numeric_non_blank_cell(self):
        """A stray formula-error string (e.g. '#REF!' -- several other sheets
        in the real workbook are full of these) must raise, not silently
        become 0 and hide a real data-quality problem."""
        with pytest.raises(ValueError):
            ex._num("#REF!")

    def test_num_allows_blank_and_numeric(self):
        assert ex._num(None) == 0.0
        assert ex._num("") == 0.0
        assert ex._num(42) == 42.0
        assert ex._num(42.5) == 42.5

    def test_malformed_production_cell_raises(self, fixture_path):
        """An error-string production cell must fail the whole load loudly."""
        wb = openpyxl.load_workbook(fixture_path)
        wb["PET "].cell(row=15, column=ex._PET_WEEK_FIRST_COL + 1, value="#REF!")
        wb.save(fixture_path)
        with pytest.raises(ValueError):
            ex.load_dataset(fixture_path, horizon_weeks=_HORIZON)

    def test_week_misalignment_between_pet_and_snp_raises(self, fixture_path):
        """If SNP Dump's date header for a column doesn't match the 'PET '
        sheet's date for that same column offset, the load must fail rather
        than silently mis-align demand to the wrong week."""
        wb = openpyxl.load_workbook(fixture_path)
        wb["SNP Dump"].cell(row=1, column=ex._SNP_WEEK_FIRST_COL, value="01.01.2099")
        wb.save(fixture_path)
        with pytest.raises(ValueError, match="[Aa]lignment"):
            ex.load_dataset(fixture_path, horizon_weeks=_HORIZON)

    def test_duplicate_snp_measure_row_raises(self, fixture_path):
        """Two 'Total' rows for the same (SKU, measure) must fail loudly --
        there is no safe way to pick which one is authoritative."""
        wb = openpyxl.load_workbook(fixture_path)
        snp = wb["SNP Dump"]
        # Duplicate row 2 (60444 / Forecast / Total) onto a fresh blank row.
        dup_row = snp.max_row + 1
        for col in range(1, snp.max_column + 1):
            snp.cell(row=dup_row, column=col, value=snp.cell(row=2, column=col).value)
        wb.save(fixture_path)
        with pytest.raises(ValueError, match="duplicate"):
            ex.load_dataset(fixture_path, horizon_weeks=_HORIZON)


# ---------------------------------------------------------------------------
# Integration: the loaded dataset must be a drop-in replacement for
# data_gen.generate() -- the service layer must run over it without error.
# ---------------------------------------------------------------------------

class TestServiceLayerCompatibility:
    def test_build_supply_runs_without_error(self, dataset):
        from backend.service import build_supply

        supply_df = build_supply(dataset)
        assert len(supply_df) == 3 * _HORIZON
        assert set(supply_df["colour"]).issubset({
            "dark_blue", "light_blue", "green", "amber", "red", "dark_red", "black",
        })

    def test_build_production_runs_without_error(self, dataset):
        from backend.service import build_production

        prod = build_production(dataset)
        assert len(prod["week_totals"]) == _HORIZON

    def test_summary_runs_without_error(self, dataset):
        from backend.service import build_supply, summary

        supply_df = build_supply(dataset)
        summ = summary(supply_df, dataset)
        assert sum(summ["counts"].values()) == 3 * _HORIZON

    def test_excel_export_runs_without_error(self, dataset):
        """The real-data path must also work through the export feature."""
        from backend.export import build_excel_export, build_pdf_export

        xlsx = build_excel_export(dataset, {})
        assert xlsx[:4] == b"PK\x03\x04"
        pdf = build_pdf_export(dataset, {})
        assert pdf[:5] == b"%PDF-"
