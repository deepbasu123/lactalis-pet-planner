"""tests/test_export.py

Tests for backend.export — the presentation-only builders. The rule engine
now lives in SQL, so these fabricate plausible Gold grids (supply_df / prod /
summ) from the synthetic dataset and feed them to the builders. No warehouse,
no engine. The HTTP /api/export path is covered by tests/test_api_live.py.
"""
import io

import openpyxl
import pandas as pd
import pytest

from backend import export as export_mod
from backend.data_gen import generate

_EXPECTED_SHEETS = [
    "Overview", "Parameters", "SKUs", "Weeks", "Summary", "Capacity Flags",
    "Supply Grid", "Supply Colour Map", "Production Plan", "Production Qty Map",
]
_CANONICAL_COLOURS = {"dark_blue", "light_blue", "green", "amber", "red", "dark_red", "black"}
_CANON = ["dark_blue", "light_blue", "green", "amber", "red", "dark_red", "black"]


def _fabricate(ds):
    """Build supply_df / prod / summ in the exact shapes the SQL engine emits."""
    sku, week, demand, plan = ds["sku"], ds["week"].sort_values("horizon_index"), ds["demand"], ds["plan_line"]
    dmap = {(r.sku_code, r.week_key): max(r.forecast, r.sales_order) for r in demand.itertuples()}
    pmap = {(r.sku_code, r.week_key): r.planned_qty for r in plan.itertuples()}
    rows, i = [], 0
    for s in sku.itertuples():
        for w in week.itertuples():
            colour = _CANON[i % 7]; i += 1
            d = float(dmap[(s.sku_code, w.week_key)]); q = float(pmap[(s.sku_code, w.week_key)])
            close = max(0.0, q - d)
            rows.append({
                "sku_code": s.sku_code, "week_key": w.week_key, "horizon_index": int(w.horizon_index),
                "opening": 0.0, "recv": q, "prod": q, "demand": d, "raw": q - d, "close": close,
                "cover_weeks": 2, "severity": _CANON.index(colour) + 1,
                "colour": colour, "display_value": close,
            })
    supply_df = pd.DataFrame(rows)
    week_totals = {w.week_key: float(sum(pmap[(s.sku_code, w.week_key)] for s in sku.itertuples())) for w in week.itertuples()}
    prod = {
        "rows": [{"sku_code": s, "week_key": wk, "planned_qty": float(q)} for (s, wk), q in pmap.items()],
        "week_totals": week_totals,
        "week_flags": {w.week_key: {"R1": False, "R2": False, "R3": False, "R4": False, "no_rule": False, "over": 0} for w in week.itertuples()},
        "changeovers": 0,
    }
    counts: dict[str, int] = {}
    for r in rows:
        counts[r["colour"]] = counts.get(r["colour"], 0) + 1
    summ = {"counts": counts, "original_vs_plan": {"original": counts, "working": counts}}
    return supply_df, prod, summ


@pytest.fixture(scope="module")
def ds():
    return generate()


@pytest.fixture(scope="module")
def grids(ds):
    return _fabricate(ds)


class TestGatherExportData:
    def test_returns_expected_keys(self, ds, grids):
        data = export_mod.gather_export_data(ds, *grids)
        assert set(data.keys()) == {"supply_df", "prod", "summ"}

    def test_supply_df_has_572_rows(self, ds, grids):
        data = export_mod.gather_export_data(ds, *grids)
        assert len(data["supply_df"]) == 572


class TestBuildExcelExportUnit:
    def test_valid_workbook_with_all_sheets(self, ds, grids):
        content = export_mod.build_excel_export(ds, *grids)
        assert content[:4] == b"PK\x03\x04"
        wb = openpyxl.load_workbook(io.BytesIO(content))
        assert wb.sheetnames == _EXPECTED_SHEETS

    def test_row_counts(self, ds, grids):
        wb = openpyxl.load_workbook(io.BytesIO(export_mod.build_excel_export(ds, *grids)))
        assert wb["SKUs"].max_row == 12       # header + 11 SKUs
        assert wb["Weeks"].max_row == 53       # header + 52 weeks
        assert wb["Supply Grid"].max_row == 573  # header + 11 x 52

    def test_survives_non_finite_supply_value(self, ds, grids):
        supply_df, prod, summ = grids
        poisoned = supply_df.copy()
        poisoned.loc[0, "close"] = float("inf")
        content = export_mod.build_excel_export(ds, poisoned, prod, summ)
        assert content[:4] == b"PK\x03\x04"


class TestBuildPdfExportUnit:
    def test_valid_pdf(self, ds, grids):
        content = export_mod.build_pdf_export(ds, *grids)
        assert content[:5] == b"%PDF-"
        assert len(content) > 4000

    def test_survives_a_very_long_note(self, ds, grids):
        d2 = {k: (v.copy() if hasattr(v, "copy") else v) for k, v in ds.items()}
        d2["week"].loc[d2["week"]["week_key"] == "2026-W38", "note"] = "x" * 100_000
        assert export_mod.build_pdf_export(d2, *grids)[:5] == b"%PDF-"


class TestColourPalette:
    def test_colour_keys_consistent(self):
        for name in export_mod.COLOUR_ORDER:
            assert name in export_mod.COLOUR_HEX
            assert name in export_mod.COLOUR_LABEL
            assert name in export_mod.COLOUR_DESC
            assert name in export_mod.COLOUR_CODE

    def test_seven_canonical_colours(self):
        assert set(export_mod.COLOUR_ORDER) == _CANONICAL_COLOURS

    def test_amber_uses_dark_foreground(self):
        assert export_mod._fg_hex("amber") == "1A2733"

    def test_non_amber_uses_light_foreground(self):
        for name in _CANONICAL_COLOURS - {"amber"}:
            assert export_mod._fg_hex(name) == "FFFFFF"


class TestWeekShort:
    def test_strips_year_prefix(self):
        assert export_mod._week_short("2026-W35") == "W35"

    def test_returns_input_unchanged_if_not_iso_week(self):
        assert export_mod._week_short("not-a-week") == "not-a-week"


class TestEditNoteAndClean:
    def test_zero_edits(self):
        assert export_mod._edit_note(0) == "Reflects the saved production plan."

    def test_clean_num_defaults_and_passthrough(self):
        assert export_mod._clean_num(float("inf")) == 0.0
        assert export_mod._clean_num(float("nan")) == 0.0
        assert export_mod._clean_num(None) == 0.0
        assert export_mod._clean_num(42.5) == 42.5

    def test_truncate_for_pdf_caps_long_text(self):
        truncated = export_mod._truncate_for_pdf("x" * 1000)
        assert len(truncated) <= export_mod._MAX_NOTE_CHARS_PDF + 1
        assert truncated.endswith("…")


class TestExcelFormulaInjection:
    @pytest.mark.parametrize("payload", ['=HYPERLINK("http://evil", "x")', "+1+1", "@SUM(1,1)"])
    def test_formula_trigger_note_is_written_as_plain_text(self, grids, payload):
        ds = generate()
        ds["week"].loc[ds["week"]["week_key"] == "2026-W38", "note"] = payload
        content = export_mod.build_excel_export(ds, *grids)
        wb = openpyxl.load_workbook(io.BytesIO(content))
        ws = wb["Weeks"]
        note_col = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))].index("note") + 1
        cell = next(row[note_col - 1] for row in ws.iter_rows(min_row=2) if row[0].value == "2026-W38")
        assert cell.data_type == "s"
        assert cell.value == "'" + payload
