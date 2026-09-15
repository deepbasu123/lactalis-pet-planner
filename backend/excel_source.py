"""backend/excel_source.py

Parses a real-world "PET Traffic Lights" workbook (the Lactalis SNP/APO
export plus the manually-maintained production plan) into the same
dict-of-DataFrames shape as backend.data_gen.generate(), so deploy.py can
load Unity Catalog with real operational data instead of the synthetic
demo dataset.

This module never reads a file on import -- call load_dataset(path)
explicitly. No workbook path is hard-coded here; the real, sensitive
workbook is never committed to the repo, only referenced by a local path
the caller supplies at deploy time (see deploy.py --data-file).

Workbook shape (verified against the 07.09.2026 export; column POSITIONS
are used throughout rather than header-text matching, since header rows
are sparse/inconsistent -- positions were confirmed stable by cross-
checking values, e.g. SNP Dump's "Stock on Hand (Projected)" Initial
column exactly matches the separate "SOH report" sheet for every SKU that
appears in both):

  "PET " sheet (trailing space is part of the real sheet name)
    Row 13, col I (9) onward  -> week_commencing dates, one column/week
    Row 12, col I (9) onward  -> "MAINT" marker for maintenance weeks
    Rows 15-25, col A/B/C     -> Status / Sku # / Sku Description
                                 (11 active SKUs, in the planner's own order)
    Rows 15-25, col I onward  -> planned production qty (EA) per week

  "SNP Dump" sheet (raw APO/SNP export: ~34 measure rows per SKU, each
  split into Total/ROW/VIC2 location sub-rows)
    Col B (2) -> SKU code (int)
    Col C (3) -> measure name (None for the base/default demand row)
    Col E (5) -> location ("Total" is the aggregate used throughout)
    Col G (7) -> "Initial" value (used only for opening stock)
    Col H (8) onward -> one column per week, positionally aligned with the
                         "PET " sheet's week columns (both start 07.09.2026)
    Measures used: "Forecast", "Sales Order", "DistrDemand (Planned)",
    "DistrDemand (TLB Confirmed)", "Stock on Hand (Projected)".

  "MLOR" sheet
    Col A(1)/C(3)/E(5) -> Product code / shelf life (days) / current MLOR (days)
    Only 6 of the 11 SKUs are covered -- see _MLOR_FALLBACK.

Data-quality notes (surfaced to the user, not silently hidden):
  - Production is a REAL, partially-populated plan: most week cells are
    blank (nothing committed that far out yet) and load as 0.0 -- not
    fabricated to look "complete".
  - MLOR/shelf-life is missing for 5 of 11 SKUs in the source file; those
    are imputed from the closest real analog SKU IN THE SAME FILE (same
    pack size, same product family) -- never an invented number. See
    _MLOR_FALLBACK.
  - orig_qty == planned_qty on load: the source file has no separate
    "baseline vs. working plan" concept. That distinction only starts to
    matter once a user edits inside the app.
  - `parameter` is NOT sourced from this workbook (capacity ceilings, QA
    hold weeks, etc. are business-rule constants, not a per-refresh data
    pull) -- it reuses backend.data_gen's existing constants unchanged.
"""
from __future__ import annotations

import datetime
import logging
import re
from typing import Any

import openpyxl
import pandas as pd

logger = logging.getLogger(__name__)

# ---------------------------------------------------------------------------
# Column/row positions in the source workbook (see module docstring)
# ---------------------------------------------------------------------------

_PET_SHEET = "PET "
_SNP_SHEET = "SNP Dump"
_MLOR_SHEET = "MLOR"

_PET_SKU_ROWS = range(15, 26)   # candidate SKU rows, planner's own order
_PET_STATUS_COL = 1
_PET_CODE_COL = 2
_PET_DESC_COL = 3
_PET_WEEK_DATE_ROW = 13
_PET_MAINT_ROW = 12
_PET_WEEK_FIRST_COL = 9          # column I
_PET_MAINT_MARKER = "MAINT"
_PET_ACTIVE_STATUS = "Active"

_SNP_CODE_COL = 2
_SNP_MEASURE_COL = 3
_SNP_LOCATION_COL = 5
_SNP_INITIAL_COL = 7
_SNP_WEEK_FIRST_COL = 8           # column H
_SNP_LOCATION_TOTAL = "Total"

_MLOR_CODE_COL = 1
_MLOR_SHELF_COL = 3    # "Manufacutre" (shelf life, days) -- sic, source's own header typo
_MLOR_MLOR_COL = 5     # "Current MLOR" (days)
_MLOR_FIRST_DATA_ROW = 4

_HORIZON_WEEKS = 52
_LOCK_WEEKS = 3   # business rule (RULE-020), independent of source data -- matches data_gen.py

# Shelf-life/MLOR is missing for these 5 SKUs in the source MLOR sheet.
# Each is imputed from the closest real analog IN THE SAME FILE (identical
# pack size + product family), never an invented number.
_MLOR_FALLBACK: dict[str, str] = {
    "228500": "60444",   # PAULS PLUS CHOC 400ml            <- PAULS ZYMIL CHOC 400ml (same class)
    "228510": "60444",   # PAULS PLUS BANANA HONEY 400ml    <- same
    "230540": "60444",   # Pauls PLUS SUMMER BERRIES 400ml  <- same
    "230550": "60444",   # Pauls PLUS DOUBLE ESPRESSO 400ml <- same
    "230150": "70526",   # OAK PLUS NAS SALTED CARAMEL 6x500ml <- OAK PLUS NAS CHOC 6x500ml (same family)
}

# demand table field -> SNP Dump measure name (Total location)
_DEMAND_MEASURES: dict[str, str] = {
    "forecast": "Forecast",
    "sales_order": "Sales Order",
    "distr_demand_planned": "DistrDemand (Planned)",
    "distr_demand_tlb": "DistrDemand (TLB Confirmed)",
}
_SOH_MEASURE = "Stock on Hand (Projected)"

_PACK_SIZE_RE = re.compile(r"(400|500)\s*ml", re.IGNORECASE)


# ---------------------------------------------------------------------------
# Small pure helpers (unit-tested directly)
# ---------------------------------------------------------------------------

def _clean_description(desc: Any) -> str:
    """Collapse internal/non-breaking whitespace and strip trailing '**' markers."""
    if desc is None:
        return ""
    text = str(desc).replace("\xa0", " ")
    text = re.sub(r"\s+", " ", text).strip()
    return text.rstrip("*").strip()


def infer_pack_size_ml(description: str) -> int:
    """Derive the pack size from the SKU description text (e.g. '... 400ML' -> 400).

    Raises ValueError if neither 400 nor 500 ml appears -- fail loudly rather
    than silently defaulting, since every SKU in this line is one or the other.
    """
    m = _PACK_SIZE_RE.search(description)
    if not m:
        raise ValueError(f"Could not infer pack size (400/500 ml) from description: {description!r}")
    return int(m.group(1))


def iso_week_key(d: datetime.date) -> str:
    """'2026-09-07' (a Monday) -> '2026-W37'."""
    y, w, _ = d.isocalendar()
    return f"{y}-W{w:02d}"


def _as_date(v: Any) -> datetime.date:
    if isinstance(v, datetime.datetime):
        return v.date()
    if isinstance(v, datetime.date):
        return v
    raise TypeError(f"Expected a date/datetime cell, got {v!r}")


def _num(v: Any) -> float:
    """Blank/None cells in a real, partially-populated plan mean 'nothing
    scheduled/reported yet' -- 0.0, not missing data to guess at.

    A cell that HAS a value but isn't numeric (e.g. a stale formula error
    like '#REF!' or '#DIV/0!' -- several other sheets in this workbook are
    full of them) is a real data-quality problem, not a blank. Silently
    treating it as 0 would hide exactly the kind of error a planner needs
    to know about, so this raises instead.
    """
    if v is None or v == "":
        return 0.0
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        return float(v)
    raise ValueError(f"Expected a blank cell or a number, got {v!r}")


def _parse_snp_date(v: Any) -> datetime.date:
    """SNP Dump stores week-column headers as 'DD.MM.YYYY' strings."""
    if isinstance(v, (datetime.datetime, datetime.date)):
        return _as_date(v)
    return datetime.datetime.strptime(str(v), "%d.%m.%Y").date()


# ---------------------------------------------------------------------------
# Sheet readers
# ---------------------------------------------------------------------------

def _read_sku_rows(pet_ws) -> list[tuple[int, str, str, int]]:
    """Return [(priority, sku_code, description, source_row), ...] for ACTIVE SKUs only.

    Priority is 1..N in the exact order the planner listed the active SKUs
    in the "PET " sheet -- used as-is rather than re-deriving an arbitrary
    order. Status is read from column A (not assumed): a row present in the
    candidate range but marked inactive/discontinued in a future refresh of
    this recurring export must not silently end up in the plan as "Active".

    source_row is carried alongside priority (rather than re-derived from it
    via _PET_SKU_ROWS[priority - 1]) because priority only counts ACTIVE
    rows: if any candidate row is skipped, priority and row position diverge,
    and re-deriving the row from priority would silently read every
    subsequent SKU's production plan off the wrong worksheet row.
    """
    rows: list[tuple[int, str, str, int]] = []
    priority = 0
    for r in _PET_SKU_ROWS:
        code = pet_ws.cell(row=r, column=_PET_CODE_COL).value
        if code is None:
            continue
        status = pet_ws.cell(row=r, column=_PET_STATUS_COL).value
        if str(status).strip().lower() != _PET_ACTIVE_STATUS.lower():
            logger.info("Skipping SKU %s: status is %r, not Active", code, status)
            continue
        priority += 1
        desc = pet_ws.cell(row=r, column=_PET_DESC_COL).value
        rows.append((priority, str(int(code)), _clean_description(desc), r))
    return rows


def _validate_week_alignment(snp_ws, week_dates: list[datetime.date]) -> None:
    """Fail loudly if SNP Dump's week columns don't positionally match PET's.

    "PET " and "SNP Dump" are independently maintained (one hand-typed, one
    pasted from an external SNP/APO export) -- a future refresh could add,
    remove, or reorder a column in one but not the other. Demand silently
    mapped to the wrong week would be far worse than a loud failure here.
    """
    for i, expected in enumerate(week_dates):
        raw = snp_ws.cell(row=1, column=_SNP_WEEK_FIRST_COL + i).value
        try:
            actual = _parse_snp_date(raw)
        except (TypeError, ValueError) as exc:
            raise ValueError(
                f"SNP Dump week column offset {i} (expected date "
                f"{expected.isoformat()}) has an unparseable header {raw!r}"
            ) from exc
        if actual != expected:
            raise ValueError(
                f"Week alignment mismatch at offset {i}: 'PET ' sheet says "
                f"{expected.isoformat()}, 'SNP Dump' says {actual.isoformat()}. "
                "Refusing to guess -- demand cannot be safely aligned to the "
                "week horizon."
            )


def _read_mlor(mlor_ws) -> dict[str, tuple[int, int]]:
    """sku_code -> (shelf_life_days, mlor_days) for every SKU present in MLOR."""
    result: dict[str, tuple[int, int]] = {}
    for r in range(_MLOR_FIRST_DATA_ROW, mlor_ws.max_row + 1):
        code = mlor_ws.cell(row=r, column=_MLOR_CODE_COL).value
        if code is None:
            continue
        shelf = mlor_ws.cell(row=r, column=_MLOR_SHELF_COL).value
        mlor_days = mlor_ws.cell(row=r, column=_MLOR_MLOR_COL).value
        if shelf is None or mlor_days is None:
            continue
        result[str(int(code))] = (int(shelf), int(mlor_days))
    return result


def _read_week_dates(pet_ws, horizon_weeks: int) -> list[datetime.date]:
    dates = []
    for i in range(horizon_weeks):
        col = _PET_WEEK_FIRST_COL + i
        v = pet_ws.cell(row=_PET_WEEK_DATE_ROW, column=col).value
        dates.append(_as_date(v))
    return dates


def _read_maintenance_offsets(pet_ws, horizon_weeks: int) -> set[int]:
    """0-indexed offsets into the week horizon flagged 'MAINT' in the source."""
    offsets: set[int] = set()
    for i in range(horizon_weeks):
        col = _PET_WEEK_FIRST_COL + i
        if pet_ws.cell(row=_PET_MAINT_ROW, column=col).value == _PET_MAINT_MARKER:
            offsets.add(i)
    return offsets


def _index_snp_rows(snp_ws, codes: list[str]) -> dict[tuple[str, Any], int]:
    """(sku_code, measure) -> row number, location == 'Total' only.

    measure is None for the base/default demand row (not currently
    consumed, indexed anyway for completeness/debuggability).

    Raises if the same (sku_code, measure) key appears more than once --
    the block structure of this export means that should never happen;
    silently keeping "whichever row came last" could quietly pick the
    wrong figure.
    """
    code_ints = {int(c) for c in codes}
    index: dict[tuple[str, Any], int] = {}
    for r in range(1, snp_ws.max_row + 1):
        code_val = snp_ws.cell(row=r, column=_SNP_CODE_COL).value
        if code_val not in code_ints:
            continue
        if snp_ws.cell(row=r, column=_SNP_LOCATION_COL).value != _SNP_LOCATION_TOTAL:
            continue
        measure = snp_ws.cell(row=r, column=_SNP_MEASURE_COL).value
        key = (str(code_val), measure)
        if key in index:
            raise ValueError(
                f"SNP Dump has duplicate {_SNP_LOCATION_TOTAL!r} rows for SKU "
                f"{code_val} / measure {measure!r} (rows {index[key]} and {r}) -- "
                "cannot safely determine which is authoritative."
            )
        index[key] = r
    return index


def _snp_week_values(snp_ws, row: int, horizon_weeks: int) -> list[float]:
    return [
        _num(snp_ws.cell(row=row, column=_SNP_WEEK_FIRST_COL + i).value)
        for i in range(horizon_weeks)
    ]


# ---------------------------------------------------------------------------
# Table builders
# ---------------------------------------------------------------------------

def _build_sku_df(
    sku_rows: list[tuple[int, str, str, int]],
    mlor_by_code: dict[str, tuple[int, int]],
) -> pd.DataFrame:
    records = []
    for priority, code, desc, _row in sku_rows:
        pack_ml = infer_pack_size_ml(desc)
        if code in mlor_by_code:
            shelf, mlor_days = mlor_by_code[code]
        else:
            fallback_code = _MLOR_FALLBACK[code]
            shelf, mlor_days = mlor_by_code[fallback_code]
            logger.info(
                "SKU %s has no MLOR sheet entry -- imputing shelf_life=%d/mlor=%d from SKU %s",
                code, shelf, mlor_days, fallback_code,
            )
        records.append({
            "sku_code": code,
            "description": desc,
            "pack_size_ml": pack_ml,
            "priority": priority,
            "status": "Active",
            "shelf_life_days": shelf,
            "mlor_days": mlor_days,
            "max_cover_weeks": round((shelf - mlor_days) / 7.0, 1),
        })
    return pd.DataFrame(records)


def _build_week_df(
    week_dates: list[datetime.date], maint_offsets: set[int]
) -> pd.DataFrame:
    records = []
    for i, d in enumerate(week_dates):
        hi = i + 1
        records.append({
            "week_key": iso_week_key(d),
            "horizon_index": hi,
            "week_commencing": d,
            "maintenance_type": "Full" if i in maint_offsets else "None",
            "is_locked": hi <= _LOCK_WEEKS,
            "note": "",
        })
    return pd.DataFrame(records)


def _build_demand_df(
    snp_ws, codes: list[str], week_keys: list[str]
) -> pd.DataFrame:
    index = _index_snp_rows(snp_ws, codes)
    horizon_weeks = len(week_keys)

    records = []
    for code in codes:
        series: dict[str, list[float]] = {}
        for field, measure in _DEMAND_MEASURES.items():
            row = index.get((code, measure))
            if row is None:
                raise KeyError(
                    f"SNP Dump has no {measure!r} / Total row for SKU {code}"
                )
            series[field] = _snp_week_values(snp_ws, row, horizon_weeks)
        for i, wk in enumerate(week_keys):
            records.append({
                "sku_code": code,
                "week_key": wk,
                "forecast": series["forecast"][i],
                "sales_order": series["sales_order"][i],
                "distr_demand_planned": series["distr_demand_planned"][i],
                "distr_demand_tlb": series["distr_demand_tlb"][i],
            })
    return pd.DataFrame(records)


def _build_opening_df(snp_ws, codes: list[str]) -> pd.DataFrame:
    index = _index_snp_rows(snp_ws, codes)
    records = []
    for code in codes:
        row = index.get((code, _SOH_MEASURE))
        if row is None:
            raise KeyError(f"SNP Dump has no {_SOH_MEASURE!r} / Total row for SKU {code}")
        initial = _num(snp_ws.cell(row=row, column=_SNP_INITIAL_COL).value)
        records.append({"sku_code": code, "opening_ea": initial})
    return pd.DataFrame(records)


def _build_plan_df(
    pet_ws, sku_rows: list[tuple[int, str, str, int]], week_keys: list[str]
) -> pd.DataFrame:
    records = []
    for _priority, code, _desc, r in sku_rows:
        for i, wk in enumerate(week_keys):
            qty = _num(pet_ws.cell(row=r, column=_PET_WEEK_FIRST_COL + i).value)
            # No separate baseline exists in the source file -- the current
            # plan IS the baseline until a user edits it inside the app.
            records.append({
                "sku_code": code, "week_key": wk, "planned_qty": qty, "orig_qty": qty,
            })
    return pd.DataFrame(records)


# ---------------------------------------------------------------------------
# Public interface
# ---------------------------------------------------------------------------

def load_dataset(path: str, horizon_weeks: int = _HORIZON_WEEKS) -> dict[str, pd.DataFrame]:
    """Parse a "PET Traffic Lights" workbook into the app's dataset shape.

    Returns the same dict-of-DataFrames shape as backend.data_gen.generate():
    sku, week, parameter, demand, plan_line, opening_stock.

    Raises a clear KeyError/ValueError (rather than silently producing wrong
    numbers) if an expected sheet, SKU, or measure row is missing -- this is
    real business data, not a demo dataset that can afford to guess.
    """
    wb = openpyxl.load_workbook(path, data_only=True)
    for sheet in (_PET_SHEET, _SNP_SHEET, _MLOR_SHEET):
        if sheet not in wb.sheetnames:
            raise KeyError(f"Workbook is missing the expected {sheet!r} sheet")

    pet = wb[_PET_SHEET]
    snp = wb[_SNP_SHEET]
    mlor = wb[_MLOR_SHEET]

    sku_rows = _read_sku_rows(pet)
    if not sku_rows:
        raise ValueError(f"No active SKUs found in {_PET_SHEET!r} rows {list(_PET_SKU_ROWS)}")

    mlor_by_code = _read_mlor(mlor)
    week_dates = _read_week_dates(pet, horizon_weeks)
    week_keys = [iso_week_key(d) for d in week_dates]
    maint_offsets = _read_maintenance_offsets(pet, horizon_weeks)
    codes = [code for _, code, _, _ in sku_rows]

    _validate_week_alignment(snp, week_dates)

    logger.info(
        "Parsed %d SKUs and %d weeks (%s to %s) from %s",
        len(sku_rows), horizon_weeks, week_keys[0], week_keys[-1], path,
    )

    # parameter is a set of business-rule constants, not a per-refresh data
    # pull -- reuse backend.data_gen's existing, already-validated defaults.
    from backend.data_gen import _PARAM_ROWS

    return {
        "sku": _build_sku_df(sku_rows, mlor_by_code),
        "week": _build_week_df(week_dates, maint_offsets),
        "parameter": pd.DataFrame(
            [{"name": n, "value": v, "description": d} for n, v, d in _PARAM_ROWS]
        ),
        "demand": _build_demand_df(snp, codes, week_keys),
        "plan_line": _build_plan_df(pet, sku_rows, week_keys),
        "opening_stock": _build_opening_df(snp, codes),
    }
