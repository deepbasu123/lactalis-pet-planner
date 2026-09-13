"""backend/export.py

Export builders for the Lactalis PET Line Planner.

Two formats, both derived from the exact same service-layer computation
used by the live API (backend.service.build_supply / build_production /
summary), so an export is always byte-consistent with what the grids show
for the same session overlay:

  build_excel_export(dataset, plan_overlay) -> bytes   (.xlsx, multi-sheet)
  build_pdf_export(dataset, plan_overlay)   -> bytes   (.pdf, formatted report)

Excel is the full raw-data dump (every table + computed grid, one sheet
each). PDF is a shorter, narrative "report": headline metrics, the
traffic-light distribution, capacity breaches, SKU/week configuration, and
a compact SKU-by-week colour heat map -- everything the app shows, laid
out for printing/sharing rather than for further data manipulation.

Colour palette mirrors frontend/src/colours.ts exactly so exported
documents match the on-screen traffic-light grid.
"""
from __future__ import annotations

import datetime
import io
import math
from typing import Any
from xml.sax.saxutils import escape as _xml_escape

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from backend import service

# ---------------------------------------------------------------------------
# Shared colour palette (mirrors frontend/src/colours.ts and TrafficLightSummary)
# ---------------------------------------------------------------------------

BRAND_BLUE = "004B85"

COLOUR_HEX: dict[str, str] = {
    "green":      "2E9E5B",
    "dark_blue":  "1565C0",
    "light_blue": "5BC5F2",
    "amber":      "E8A317",
    "red":        "E1382D",
    "dark_red":   "8B1A1A",
    "black":      "2B2B2B",
}

COLOUR_LABEL: dict[str, str] = {
    "green":      "Green",
    "dark_blue":  "Dark Blue",
    "light_blue": "Light Blue",
    "amber":      "Amber",
    "red":        "Red",
    "dark_red":   "Dark Red",
    "black":      "Black",
}

COLOUR_DESC: dict[str, str] = {
    "green":      "Cover 1-2 weeks",
    "dark_blue":  "Cover > 3 weeks",
    "light_blue": "Cover 2-3 weeks",
    "amber":      "Cover 0-1 weeks",
    "red":        "Stock-out (recoverable)",
    "dark_red":   "Lost sale risk",
    "black":      "Over-cover (past MLOR)",
}

COLOUR_CODE: dict[str, str] = {
    "green": "GN", "dark_blue": "DB", "light_blue": "LB",
    "amber": "AM", "red": "RD", "dark_red": "DR", "black": "BK",
}

# Same order the frontend TrafficLightSummary tab uses (its BANDS array).
COLOUR_ORDER: list[str] = [
    "green", "dark_blue", "light_blue", "amber", "red", "dark_red", "black",
]

_QTY_FORMAT = "#,##0"
_QTY_COLUMNS = {
    "opening", "recv", "prod", "demand", "raw", "close", "display_value",
    "planned_qty", "orig_qty", "total_ea", "units_over_ceiling",
}


def _fg_hex(colour_name: str) -> str:
    """Foreground hex for a colour cell -- amber needs dark text for contrast."""
    return "1A2733" if colour_name == "amber" else "FFFFFF"


def _clean(v: Any) -> Any:
    """Defensive NaN/inf -> None; numpy scalar -> python scalar."""
    if v is None:
        return None
    if hasattr(v, "item"):
        v = v.item()
    if isinstance(v, float) and (math.isnan(v) or math.isinf(v)):
        return None
    return v


def _week_short(week_key: str) -> str:
    """'2026-W35' -> 'W35'; falls back to the raw key if the shape is unexpected."""
    if "-W" in week_key:
        return "W" + week_key.split("-W")[-1]
    return week_key


def _sku_order(sku_df: pd.DataFrame) -> list[str]:
    return list(sku_df.sort_values("priority")["sku_code"])


def gather_export_data(dataset: dict, plan_overlay: dict | None = None) -> dict[str, Any]:
    """Run the service layer once; return everything both exporters need."""
    plan_overlay = plan_overlay or {}
    supply_df = service.build_supply(dataset, plan_overlay=plan_overlay)
    prod = service.build_production(dataset, plan_overlay=plan_overlay)
    summ = service.summary(supply_df, dataset, plan_overlay=plan_overlay)
    return {"supply_df": supply_df, "prod": prod, "summ": summ}


# ===========================================================================
# Excel export
# ===========================================================================

def _style_header_row(ws: Worksheet, n_cols: int, row: int = 1) -> None:
    fill = PatternFill(start_color=BRAND_BLUE, end_color=BRAND_BLUE, fill_type="solid")
    font = Font(color="FFFFFF", bold=True)
    for col in range(1, n_cols + 1):
        cell = ws.cell(row=row, column=col)
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.freeze_panes = ws.cell(row=row + 1, column=1).coordinate


def _autosize(ws: Worksheet, df: pd.DataFrame, min_width: int = 8, max_width: int = 42) -> None:
    for i, col in enumerate(df.columns, start=1):
        series_len = int(df[col].astype(str).map(len).max()) if len(df) else 0
        width = max(min_width, min(max_width, max(len(str(col)), series_len) + 2))
        ws.column_dimensions[get_column_letter(i)].width = width


def _apply_number_formats(ws: Worksheet, df: pd.DataFrame, start_row: int = 2) -> None:
    for i, col in enumerate(df.columns, start=1):
        if col in _QTY_COLUMNS:
            fmt = _QTY_FORMAT
        elif col == "max_cover_weeks":
            fmt = "0.0"
        elif col == "cover_weeks":
            fmt = "0"
        else:
            continue
        for r in range(start_row, start_row + len(df)):
            ws.cell(row=r, column=i).number_format = fmt


def _write_flat_sheet(writer: pd.ExcelWriter, sheet_name: str, df: pd.DataFrame) -> Worksheet:
    df.to_excel(writer, sheet_name=sheet_name, index=False)
    ws = writer.sheets[sheet_name]
    _style_header_row(ws, len(df.columns))
    _autosize(ws, df)
    _apply_number_formats(ws, df)
    return ws


def _build_capacity_flags_df(prod: dict, week_df: pd.DataFrame) -> pd.DataFrame:
    week_order = list(week_df.sort_values("horizon_index")["week_key"])
    commencing = dict(zip(week_df["week_key"], week_df["week_commencing"]))
    rows = []
    for wk in week_order:
        f = prod["week_flags"].get(wk, {})
        rows.append({
            "week_key": wk,
            "week_commencing": commencing.get(wk),
            "total_ea": _clean(prod["week_totals"].get(wk, 0.0)),
            "R1_multi_pack": bool(f.get("R1", False)),
            "R2_over_3_skus": bool(f.get("R2", False)),
            "R3_over_ceiling": bool(f.get("R3", False)),
            "R4_full_maint_production": bool(f.get("R4", False)),
            "no_rule_determined": bool(f.get("no_rule", False)),
            "units_over_ceiling": _clean(f.get("over", 0)),
        })
    return pd.DataFrame(rows)


def _build_supply_flat_df(supply_df: pd.DataFrame, sku_df: pd.DataFrame) -> pd.DataFrame:
    desc_map = dict(zip(sku_df["sku_code"], sku_df["description"]))
    df = supply_df.copy()
    df.insert(1, "description", df["sku_code"].map(desc_map))
    cols = [
        "sku_code", "description", "week_key", "horizon_index", "opening",
        "recv", "prod", "demand", "raw", "close", "cover_weeks", "severity",
        "colour", "display_value",
    ]
    return df[cols].sort_values(["sku_code", "horizon_index"]).reset_index(drop=True)


def _build_production_flat_df(prod: dict, dataset: dict, sku_df: pd.DataFrame) -> pd.DataFrame:
    desc_map = dict(zip(sku_df["sku_code"], sku_df["description"]))
    orig_map = {
        (r["sku_code"], r["week_key"]): r["orig_qty"]
        for r in dataset["plan_line"][["sku_code", "week_key", "orig_qty"]].to_dict("records")
    }
    week_order = {
        wk: i for i, wk in enumerate(dataset["week"].sort_values("horizon_index")["week_key"])
    }
    rows = []
    for r in prod["rows"]:
        key = (r["sku_code"], r["week_key"])
        rows.append({
            "sku_code":    r["sku_code"],
            "description": desc_map.get(r["sku_code"], ""),
            "week_key":    r["week_key"],
            "planned_qty": _clean(r["planned_qty"]),
            "orig_qty":    _clean(orig_map.get(key, r["planned_qty"])),
        })
    df = pd.DataFrame(rows)
    df["_wi"] = df["week_key"].map(week_order)
    df = df.sort_values(["sku_code", "_wi"]).drop(columns="_wi").reset_index(drop=True)
    return df


def _build_colour_pivot(
    supply_df: pd.DataFrame, sku_df: pd.DataFrame, week_df: pd.DataFrame
) -> pd.DataFrame:
    desc_map = dict(zip(sku_df["sku_code"], sku_df["description"]))
    week_order = list(week_df.sort_values("horizon_index")["week_key"])
    sku_order = _sku_order(sku_df)

    pivot = supply_df.pivot(index="sku_code", columns="week_key", values="colour")
    pivot = pivot.reindex(index=sku_order, columns=week_order)
    pivot.insert(0, "SKU", [f"{code} - {desc_map.get(code, '')}" for code in pivot.index])
    pivot.columns = ["SKU"] + [_week_short(wk) for wk in week_order]
    return pivot.reset_index(drop=True)


def _apply_colour_map_styles(
    ws: Worksheet, n_rows: int, n_week_cols: int, data_start_row: int = 2, data_start_col: int = 2
) -> None:
    for r in range(n_rows):
        for c in range(n_week_cols):
            cell = ws.cell(row=data_start_row + r, column=data_start_col + c)
            colour_name = cell.value
            if colour_name in COLOUR_HEX:
                hexv = COLOUR_HEX[colour_name]
                cell.fill = PatternFill(start_color=hexv, end_color=hexv, fill_type="solid")
                cell.font = Font(color=_fg_hex(colour_name), bold=True, size=9)
                cell.alignment = Alignment(horizontal="center")
                cell.value = COLOUR_CODE[colour_name]


def _write_colour_legend(ws: Worksheet, start_row: int) -> None:
    ws.cell(row=start_row, column=1, value="Legend").font = Font(bold=True)
    for i, name in enumerate(COLOUR_ORDER):
        r = start_row + 1 + i
        code_cell = ws.cell(row=r, column=1, value=COLOUR_CODE[name])
        hexv = COLOUR_HEX[name]
        code_cell.fill = PatternFill(start_color=hexv, end_color=hexv, fill_type="solid")
        code_cell.font = Font(color=_fg_hex(name), bold=True)
        code_cell.alignment = Alignment(horizontal="center")
        ws.cell(row=r, column=2, value=COLOUR_LABEL[name])
        ws.cell(row=r, column=3, value=COLOUR_DESC[name])


def _write_colour_map_sheet(
    writer: pd.ExcelWriter, supply_df: pd.DataFrame, sku_df: pd.DataFrame, week_df: pd.DataFrame
) -> None:
    pivot = _build_colour_pivot(supply_df, sku_df, week_df)
    sheet_name = "Supply Colour Map"
    pivot.to_excel(writer, sheet_name=sheet_name, index=False)
    ws = writer.sheets[sheet_name]
    n_rows, n_cols = len(pivot), len(pivot.columns)

    _style_header_row(ws, n_cols)
    ws.column_dimensions["A"].width = 44
    for col_idx in range(2, n_cols + 1):
        ws.column_dimensions[get_column_letter(col_idx)].width = 4.5

    _apply_colour_map_styles(ws, n_rows, n_cols - 1, data_start_row=2, data_start_col=2)
    _write_colour_legend(ws, start_row=n_rows + 4)


def _build_qty_pivot(
    prod: dict, sku_df: pd.DataFrame, week_df: pd.DataFrame
) -> tuple[pd.DataFrame, list[str]]:
    desc_map = dict(zip(sku_df["sku_code"], sku_df["description"]))
    week_order = list(week_df.sort_values("horizon_index")["week_key"])
    sku_order = _sku_order(sku_df)

    rows_df = pd.DataFrame(prod["rows"])
    pivot = rows_df.pivot(index="sku_code", columns="week_key", values="planned_qty")
    pivot = pivot.reindex(index=sku_order, columns=week_order).fillna(0.0)

    total_row = pd.Series(
        [prod["week_totals"].get(wk, 0.0) for wk in week_order],
        index=week_order, name="__TOTAL__",
    )
    pivot = pd.concat([pivot, total_row.to_frame().T])

    pivot.insert(0, "SKU", [
        "TOTAL UNITS PRODUCED" if idx == "__TOTAL__" else f"{idx} - {desc_map.get(idx, '')}"
        for idx in pivot.index
    ])
    pivot.columns = ["SKU"] + [_week_short(wk) for wk in week_order]
    return pivot.reset_index(drop=True), week_order


def _apply_breach_header_shading(
    ws: Worksheet, week_order: list[str], week_flags: dict, header_row: int = 1, data_start_col: int = 2
) -> None:
    breach_fill = PatternFill(start_color="FDEDEC", end_color="FDEDEC", fill_type="solid")
    breach_font = Font(color="C0392B", bold=True)
    for i, wk in enumerate(week_order):
        flags = week_flags.get(wk, {})
        if any(flags.get(k) for k in ("R1", "R2", "R3", "R4")):
            cell = ws.cell(row=header_row, column=data_start_col + i)
            cell.fill = breach_fill
            cell.font = breach_font


def _write_qty_map_sheet(
    writer: pd.ExcelWriter, prod: dict, sku_df: pd.DataFrame, week_df: pd.DataFrame
) -> None:
    pivot, week_order = _build_qty_pivot(prod, sku_df, week_df)
    sheet_name = "Production Qty Map"
    pivot.to_excel(writer, sheet_name=sheet_name, index=False)
    ws = writer.sheets[sheet_name]
    n_rows, n_cols = len(pivot), len(pivot.columns)

    _style_header_row(ws, n_cols)
    ws.column_dimensions["A"].width = 44
    for col_idx in range(2, n_cols + 1):
        ws.column_dimensions[get_column_letter(col_idx)].width = 9

    for r in range(2, n_rows + 2):
        for c in range(2, n_cols + 1):
            ws.cell(row=r, column=c).number_format = _QTY_FORMAT

    total_row_idx = n_rows + 1  # +1 for the header row
    fill = PatternFill(start_color="EDF1F7", end_color="EDF1F7", fill_type="solid")
    for c in range(1, n_cols + 1):
        cell = ws.cell(row=total_row_idx, column=c)
        cell.font = Font(bold=True)
        cell.fill = fill

    _apply_breach_header_shading(ws, week_order, prod["week_flags"])


def _write_overview_sheet(
    writer: pd.ExcelWriter, dataset: dict, prod: dict, plan_overlay: dict
) -> None:
    total_production = float(sum(prod["week_totals"].values()))
    generated_at = datetime.datetime.utcnow().strftime("%Y-%m-%d %H:%M UTC")
    rows = [
        ("Report",                                 "Lactalis PET Line Planner -- Full Data Export"),
        ("Generated At",                            generated_at),
        ("SKU Count",                               int(len(dataset["sku"]))),
        ("Week Count",                              int(len(dataset["week"]))),
        ("Total Production (working plan, EA)",     round(total_production)),
        ("Pack Changeovers (horizon)",               int(prod["changeovers"])),
        ("Unsaved Session Edits Included",           int(len(plan_overlay))),
    ]
    df = pd.DataFrame(rows, columns=["Metric", "Value"])
    df.to_excel(writer, sheet_name="Overview", index=False)
    ws = writer.sheets["Overview"]
    _style_header_row(ws, 2)
    _autosize(ws, df, min_width=14, max_width=60)


def _write_summary_sheet(writer: pd.ExcelWriter, summ: dict) -> None:
    counts = summ["counts"]
    orig = summ["original_vs_plan"]["original"]
    work = summ["original_vs_plan"]["working"]

    records = []
    for name in COLOUR_ORDER:
        c, o, w = counts.get(name, 0), orig.get(name, 0), work.get(name, 0)
        records.append({
            "colour": COLOUR_LABEL[name],
            "meaning": COLOUR_DESC[name],
            "current_count": int(c),
            "original_snp_count": int(o),
            "working_plan_count": int(w),
            "change_vs_snp": int(w - o),
        })
    records.append({
        "colour": "Total", "meaning": "",
        "current_count": int(sum(counts.values())),
        "original_snp_count": int(sum(orig.values())),
        "working_plan_count": int(sum(work.values())),
        "change_vs_snp": int(sum(work.values()) - sum(orig.values())),
    })
    df = pd.DataFrame(records)
    df.to_excel(writer, sheet_name="Summary", index=False)
    ws = writer.sheets["Summary"]
    _style_header_row(ws, len(df.columns))
    _autosize(ws, df)

    for i, name in enumerate(COLOUR_ORDER, start=2):  # row 1 is the header
        hexv = COLOUR_HEX[name]
        cell = ws.cell(row=i, column=1)
        cell.fill = PatternFill(start_color=hexv, end_color=hexv, fill_type="solid")
        cell.font = Font(color=_fg_hex(name), bold=True)

    total_row = len(COLOUR_ORDER) + 2
    fill = PatternFill(start_color="EDF1F7", end_color="EDF1F7", fill_type="solid")
    for col in range(1, len(df.columns) + 1):
        cell = ws.cell(row=total_row, column=col)
        cell.font = Font(bold=True)
        cell.fill = fill


def build_excel_export(dataset: dict, plan_overlay: dict | None = None) -> bytes:
    """Build the full multi-sheet Excel workbook and return it as bytes.

    Sheets: Overview, Parameters, SKUs, Weeks, Summary, Capacity Flags,
    Supply Grid, Supply Colour Map, Production Plan, Production Qty Map.
    """
    plan_overlay = plan_overlay or {}
    data = gather_export_data(dataset, plan_overlay)
    supply_df, prod, summ = data["supply_df"], data["prod"], data["summ"]

    sku_df = dataset["sku"].sort_values("priority").reset_index(drop=True)
    week_df = dataset["week"].sort_values("horizon_index").reset_index(drop=True)
    param_df = dataset["parameter"].reset_index(drop=True)

    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as writer:
        _write_overview_sheet(writer, dataset, prod, plan_overlay)
        _write_flat_sheet(writer, "Parameters", param_df)
        _write_flat_sheet(writer, "SKUs", sku_df)
        _write_flat_sheet(writer, "Weeks", week_df)
        _write_summary_sheet(writer, summ)
        _write_flat_sheet(writer, "Capacity Flags", _build_capacity_flags_df(prod, week_df))
        _write_flat_sheet(writer, "Supply Grid", _build_supply_flat_df(supply_df, sku_df))
        _write_colour_map_sheet(writer, supply_df, sku_df, week_df)
        _write_flat_sheet(writer, "Production Plan", _build_production_flat_df(prod, dataset, sku_df))
        _write_qty_map_sheet(writer, prod, sku_df, week_df)

    return buf.getvalue()


# ===========================================================================
# PDF export
# ===========================================================================

def _p(text: Any, style) -> Any:
    """Wrap free-form text in an XML-escaped Paragraph so long values wrap."""
    from reportlab.platypus import Paragraph
    s = "" if text is None else str(text)
    return Paragraph(_xml_escape(s), style)


def _styled_table(
    rows: list[list],
    col_widths: list[float] | None = None,
    font_size: float = 8,
    header: bool = True,
    cell_styles: list[tuple] | None = None,
    padding: float = 3,
):
    from reportlab.lib import colors
    from reportlab.platypus import Table, TableStyle

    table = Table(rows, colWidths=col_widths, repeatRows=1 if header else 0)
    cmds = [
        ("FONTNAME", (0, 0), (-1, -1), "Helvetica"),
        ("FONTSIZE", (0, 0), (-1, -1), font_size),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#E2E8F0")),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("LEFTPADDING", (0, 0), (-1, -1), padding),
        ("RIGHTPADDING", (0, 0), (-1, -1), padding),
        ("TOPPADDING", (0, 0), (-1, -1), max(1.0, padding - 1)),
        ("BOTTOMPADDING", (0, 0), (-1, -1), max(1.0, padding - 1)),
    ]
    body_start = 0
    if header:
        cmds += [
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(f"#{BRAND_BLUE}")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ]
        body_start = 1
    cmds.append(("ROWBACKGROUNDS", (0, body_start), (-1, -1), [colors.white, colors.HexColor("#F5F7FA")]))
    if cell_styles:
        cmds.extend(cell_styles)
    table.setStyle(TableStyle(cmds))
    return table


def _colour_summary_table(summ: dict, usable_width: float):
    from reportlab.lib import colors

    counts = summ["counts"]
    orig = summ["original_vs_plan"]["original"]
    work = summ["original_vs_plan"]["working"]

    rows = [["Status", "Meaning", "Current", "Original (SNP)", "Working Plan", "Change"]]
    cmds = []
    for i, name in enumerate(COLOUR_ORDER, start=1):
        c, o, w = counts.get(name, 0), orig.get(name, 0), work.get(name, 0)
        delta = w - o
        rows.append([
            COLOUR_LABEL[name], COLOUR_DESC[name], str(c), str(o), str(w),
            f"+{delta}" if delta > 0 else str(delta),
        ])
        cmds.append(("BACKGROUND", (0, i), (0, i), colors.HexColor(f"#{COLOUR_HEX[name]}")))
        cmds.append(("TEXTCOLOR", (0, i), (0, i), colors.HexColor(f"#{_fg_hex(name)}")))
        cmds.append(("FONTNAME", (0, i), (0, i), "Helvetica-Bold"))

    total_delta = sum(work.values()) - sum(orig.values())
    total_row_idx = len(rows)
    rows.append([
        "Total", "", str(sum(counts.values())), str(sum(orig.values())), str(sum(work.values())),
        f"+{total_delta}" if total_delta > 0 else str(total_delta),
    ])
    cmds.append(("FONTNAME", (0, total_row_idx), (-1, total_row_idx), "Helvetica-Bold"))
    cmds.append(("BACKGROUND", (0, total_row_idx), (-1, total_row_idx), colors.HexColor("#EDF1F7")))

    col_widths = [usable_width * w for w in (0.14, 0.32, 0.13, 0.15, 0.14, 0.12)]
    return _styled_table(rows, col_widths=col_widths, font_size=8, cell_styles=cmds)


def _capacity_breach_table(prod: dict, week_df: pd.DataFrame, usable_width: float):
    week_order = list(week_df.sort_values("horizon_index")["week_key"])
    rows = [["Week", "Total EA", "Units over ceiling", "Rules breached"]]
    any_breach = False
    for wk in week_order:
        f = prod["week_flags"].get(wk, {})
        breached = [
            label for label, key in (
                ("R1 multi-pack", "R1"), ("R2 >3 SKUs", "R2"),
                ("R3 over ceiling", "R3"), ("R4 full-maint production", "R4"),
            )
            if f.get(key)
        ]
        if not breached:
            continue
        any_breach = True
        total = prod["week_totals"].get(wk, 0.0)
        over = f.get("over", 0)
        rows.append([
            _week_short(wk), f"{total:,.0f}",
            f"{over:,.0f}" if f.get("R3") else "-",
            ", ".join(breached),
        ])
    if not any_breach:
        rows.append(["--", "--", "--", "No capacity-rule breaches detected"])

    col_widths = [usable_width * w for w in (0.12, 0.16, 0.20, 0.52)]
    return _styled_table(rows, col_widths=col_widths, font_size=8)


def _sku_table(sku_df: pd.DataFrame, usable_width: float):
    rows = [["Code", "Description", "Pack (ml)", "Priority", "Status", "Shelf life (d)", "MLOR (d)", "Max cover (wks)"]]
    for _, r in sku_df.sort_values("priority").iterrows():
        rows.append([
            str(r["sku_code"]), str(r["description"]), str(int(r["pack_size_ml"])),
            str(int(r["priority"])), str(r["status"]), str(int(r["shelf_life_days"])),
            str(int(r["mlor_days"])), f"{float(r['max_cover_weeks']):.1f}",
        ])
    col_widths = [usable_width * w for w in (0.09, 0.37, 0.09, 0.08, 0.09, 0.11, 0.08, 0.09)]
    return _styled_table(rows, col_widths=col_widths, font_size=8)


def _week_table(week_df: pd.DataFrame, usable_width: float):
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle

    note_style = ParagraphStyle("PETNote", parent=getSampleStyleSheet()["BodyText"], fontSize=7.5, leading=9)
    rows = [["Week", "Commencing", "Maintenance", "Locked", "Note"]]
    for _, r in week_df.sort_values("horizon_index").iterrows():
        wc = r["week_commencing"]
        wc_str = wc.isoformat() if hasattr(wc, "isoformat") else str(wc)
        rows.append([
            _week_short(r["week_key"]), wc_str, str(r["maintenance_type"]),
            "Yes" if bool(r["is_locked"]) else "No",
            _p(r["note"] or "", note_style),  # free-text field -- wrap + XML-escape
        ])
    col_widths = [usable_width * w for w in (0.09, 0.16, 0.16, 0.09, 0.50)]
    return _styled_table(rows, col_widths=col_widths, font_size=7.5)


def _heatmap_half_table(
    supply_df: pd.DataFrame, sku_df: pd.DataFrame, week_keys_half: list[str], usable_width: float
):
    from reportlab.lib import colors

    sku_order = _sku_order(sku_df)
    colour_lookup = {
        (r["sku_code"], r["week_key"]): r["colour"]
        for r in supply_df[["sku_code", "week_key", "colour"]].to_dict("records")
    }

    header = ["SKU"] + [_week_short(wk) for wk in week_keys_half]
    rows = [header]
    cmds = [("FONTNAME", (0, 1), (0, -1), "Helvetica-Bold"), ("ALIGN", (1, 1), (-1, -1), "CENTER")]
    for ri, code in enumerate(sku_order, start=1):
        row = [code]
        for ci, wk in enumerate(week_keys_half, start=1):
            colour_name = colour_lookup.get((code, wk), "")
            row.append(COLOUR_CODE.get(colour_name, ""))
            if colour_name in COLOUR_HEX:
                cmds.append(("BACKGROUND", (ci, ri), (ci, ri), colors.HexColor(f"#{COLOUR_HEX[colour_name]}")))
                cmds.append(("TEXTCOLOR", (ci, ri), (ci, ri), colors.HexColor(f"#{_fg_hex(colour_name)}")))
        rows.append(row)

    n_cols = len(header)
    sku_col_w = 34.0
    week_col_w = (usable_width - sku_col_w) / (n_cols - 1)
    col_widths = [sku_col_w] + [week_col_w] * (n_cols - 1)
    return _styled_table(rows, col_widths=col_widths, font_size=5.5, cell_styles=cmds, padding=1.5)


def _colour_legend_table(usable_width: float):
    from reportlab.lib import colors

    rows = [["Code", "Colour", "Meaning"]]
    cmds = []
    for i, name in enumerate(COLOUR_ORDER, start=1):
        rows.append([COLOUR_CODE[name], COLOUR_LABEL[name], COLOUR_DESC[name]])
        cmds.append(("BACKGROUND", (0, i), (0, i), colors.HexColor(f"#{COLOUR_HEX[name]}")))
        cmds.append(("TEXTCOLOR", (0, i), (0, i), colors.HexColor(f"#{_fg_hex(name)}")))
        cmds.append(("FONTNAME", (0, i), (0, i), "Helvetica-Bold"))
    col_widths = [usable_width * w for w in (0.16, 0.24, 0.60)]
    return _styled_table(rows, col_widths=col_widths, font_size=8, cell_styles=cmds)


def build_pdf_export(dataset: dict, plan_overlay: dict | None = None) -> bytes:
    """Build the formatted PDF planning report and return it as bytes.

    Sections: overview, parameters, traffic-light distribution, capacity
    breaches, SKU master, week configuration, and a two-page SKU-by-week
    traffic-light heat map (weeks 1-26 / 27-52) with a colour legend.
    """
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer

    plan_overlay = plan_overlay or {}
    data = gather_export_data(dataset, plan_overlay)
    supply_df, prod, summ = data["supply_df"], data["prod"], data["summ"]

    sku_df = dataset["sku"]
    week_df = dataset["week"].sort_values("horizon_index").reset_index(drop=True)
    week_order = list(week_df["week_key"])
    mid = (len(week_order) + 1) // 2  # 52 -> 26/26; odd horizons split as evenly as possible

    buf = io.BytesIO()
    page_size = landscape(A4)
    doc = SimpleDocTemplate(
        buf, pagesize=page_size,
        leftMargin=24, rightMargin=24, topMargin=26, bottomMargin=24,
        title="Lactalis PET Line Planner Report",
    )
    usable_width = page_size[0] - doc.leftMargin - doc.rightMargin

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "PETTitle", parent=styles["Title"], textColor=colors.HexColor(f"#{BRAND_BLUE}"),
        fontSize=22, spaceAfter=2,
    )
    sub_style = ParagraphStyle(
        "PETSub", parent=styles["Heading3"], textColor=colors.HexColor("#4A5568"),
        fontSize=12, spaceAfter=2,
    )
    h2_style = ParagraphStyle(
        "PETH2", parent=styles["Heading2"], textColor=colors.HexColor(f"#{BRAND_BLUE}"),
        fontSize=13, spaceBefore=14, spaceAfter=6,
    )
    muted_style = ParagraphStyle(
        "PETMuted", parent=styles["BodyText"], textColor=colors.HexColor("#718096"), fontSize=9,
    )

    generated_at = datetime.datetime.utcnow().strftime("%Y-%m-%d %H:%M UTC")
    n_edits = len(plan_overlay)
    edit_note = (
        f"Includes {n_edits} unsaved session edit(s)." if n_edits
        else "Reflects the saved production plan."
    )
    total_production = float(sum(prod["week_totals"].values()))

    story = [
        Paragraph("Lactalis PET Line Planner", title_style),
        Paragraph("PET Line Planning Report", sub_style),
        Paragraph(f"Generated {generated_at} \u2014 {edit_note}", muted_style),
        Spacer(1, 8),

        Paragraph("Overview", h2_style),
        _styled_table(
            [
                ["Metric", "Value"],
                ["SKUs in scope", str(len(sku_df))],
                ["Planning horizon", f"{len(week_df)} weeks ({week_order[0]} to {week_order[-1]})"],
                ["Total production (working plan)", f"{total_production:,.0f} EA"],
                ["Pack changeovers across horizon", str(prod["changeovers"])],
            ],
            col_widths=[usable_width * 0.4, usable_width * 0.6], font_size=9,
        ),

        Paragraph("Planning Parameters", h2_style),
        _styled_table(
            [["Parameter", "Value", "Description"]] + [
                [str(r["name"]), f"{float(r['value']):,.1f}", str(r["description"])]
                for _, r in dataset["parameter"].iterrows()
            ],
            col_widths=[usable_width * 0.22, usable_width * 0.12, usable_width * 0.66], font_size=8,
        ),

        Paragraph("Traffic-Light Distribution", h2_style),
        _colour_summary_table(summ, usable_width),

        Paragraph("Capacity Rule Breaches", h2_style),
        _capacity_breach_table(prod, week_df, usable_width),

        PageBreak(),
        Paragraph("SKU Master", h2_style),
        _sku_table(sku_df, usable_width),

        Paragraph("Week Configuration", h2_style),
        _week_table(week_df, usable_width),

        PageBreak(),
        Paragraph(f"Supply Traffic-Light Map \u2014 Weeks 1-{mid}", h2_style),
        _heatmap_half_table(supply_df, sku_df, week_order[:mid], usable_width),
        Spacer(1, 10),
        _colour_legend_table(usable_width * 0.55),

        PageBreak(),
        Paragraph(f"Supply Traffic-Light Map \u2014 Weeks {mid + 1}-{len(week_order)}", h2_style),
        _heatmap_half_table(supply_df, sku_df, week_order[mid:], usable_width),
    ]

    doc.build(story)
    return buf.getvalue()
