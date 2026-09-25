"""medallion/synth_workbook.py

Writes a synthetic "PET Traffic Lights" workbook (.xlsx) in the EXACT 3-sheet
layout that medallion.parse_workbook reads. This lets the demo seed flow through
the identical Bronze->Silver->Gold pipeline path as a real customer upload — no
special-case seeding, one code path.

The numbers come from backend.data_gen.generate(seed), so a round-trip
(write_synth_workbook -> parse_workbook.load_dataset) reproduces the same
dataset the app was validated against.

Note: this is a seed-time tool (run locally or in a seed job), so it MAY import
backend.data_gen. The pipeline compute never imports this module.
"""
from __future__ import annotations

import openpyxl

from backend import data_gen
from medallion.parse_workbook import (
    _PET_SHEET,
    _SNP_SHEET,
    _MLOR_SHEET,
    _PET_SKU_ROWS,
    _PET_STATUS_COL,
    _PET_CODE_COL,
    _PET_DESC_COL,
    _PET_WEEK_DATE_ROW,
    _PET_MAINT_ROW,
    _PET_WEEK_FIRST_COL,
    _PET_MAINT_MARKER,
    _SNP_CODE_COL,
    _SNP_MEASURE_COL,
    _SNP_LOCATION_COL,
    _SNP_INITIAL_COL,
    _SNP_WEEK_FIRST_COL,
    _SNP_LOCATION_TOTAL,
    _MLOR_CODE_COL,
    _MLOR_SHELF_COL,
    _MLOR_MLOR_COL,
    _MLOR_FIRST_DATA_ROW,
    _DEMAND_MEASURES,
    _SOH_MEASURE,
)

_SKU_ROW_START = _PET_SKU_ROWS.start  # 15


def write_synth_workbook(path: str, seed: int = 42) -> str:
    """Generate a synthetic PET workbook at ``path`` and return the path."""
    ds = data_gen.generate(seed)
    sku = ds["sku"]
    week = ds["week"].sort_values("horizon_index").reset_index(drop=True)
    demand = ds["demand"]
    plan = ds["plan_line"]
    opening = ds["opening_stock"]

    week_keys = list(week["week_key"])
    dates = list(week["week_commencing"])
    maint = list(week["maintenance_type"])
    codes = list(sku["sku_code"])

    wb = openpyxl.Workbook()

    # ---- "PET " sheet ------------------------------------------------------
    pet = wb.active
    pet.title = _PET_SHEET
    for i, (d, m) in enumerate(zip(dates, maint)):
        col = _PET_WEEK_FIRST_COL + i
        pet.cell(row=_PET_WEEK_DATE_ROW, column=col, value=d)
        # Parser only recognises Full-maintenance MAINT markers; a Partial week
        # is carried purely by its (already reduced) production quantities.
        if m == "Full":
            pet.cell(row=_PET_MAINT_ROW, column=col, value=_PET_MAINT_MARKER)

    plan_piv = plan.pivot(index="sku_code", columns="week_key", values="planned_qty")
    sku_by_code = {row["sku_code"]: row for _, row in sku.iterrows()}
    for r, code in enumerate(codes):
        row = _SKU_ROW_START + r
        srow = sku_by_code[code]
        pet.cell(row=row, column=_PET_STATUS_COL, value="Active")
        pet.cell(row=row, column=_PET_CODE_COL, value=int(code))
        pet.cell(row=row, column=_PET_DESC_COL, value=str(srow["description"]))
        for i, wk in enumerate(week_keys):
            pet.cell(row=row, column=_PET_WEEK_FIRST_COL + i, value=float(plan_piv.loc[code, wk]))

    # ---- "SNP Dump" sheet --------------------------------------------------
    snp = wb.create_sheet(_SNP_SHEET)
    for i, d in enumerate(dates):
        snp.cell(row=1, column=_SNP_WEEK_FIRST_COL + i, value=d.strftime("%d.%m.%Y"))

    demand_by_key = {(r.sku_code, r.week_key): r for r in demand.itertuples()}
    opening_by_code = dict(zip(opening["sku_code"], opening["opening_ea"]))

    r = 2
    for code in codes:
        for field, measure_name in _DEMAND_MEASURES.items():
            snp.cell(row=r, column=_SNP_CODE_COL, value=int(code))
            snp.cell(row=r, column=_SNP_MEASURE_COL, value=measure_name)
            snp.cell(row=r, column=_SNP_LOCATION_COL, value=_SNP_LOCATION_TOTAL)
            for i, wk in enumerate(week_keys):
                snp.cell(row=r, column=_SNP_WEEK_FIRST_COL + i,
                         value=float(getattr(demand_by_key[(code, wk)], field)))
            r += 1
        # Stock on Hand (Projected): only the "Initial" column (G) is read, for
        # opening stock.
        snp.cell(row=r, column=_SNP_CODE_COL, value=int(code))
        snp.cell(row=r, column=_SNP_MEASURE_COL, value=_SOH_MEASURE)
        snp.cell(row=r, column=_SNP_LOCATION_COL, value=_SNP_LOCATION_TOTAL)
        snp.cell(row=r, column=_SNP_INITIAL_COL, value=float(opening_by_code[code]))
        r += 1

    # ---- "MLOR" sheet ------------------------------------------------------
    mlor = wb.create_sheet(_MLOR_SHEET)
    mlor.cell(row=1, column=_MLOR_CODE_COL, value="Product")
    mlor.cell(row=1, column=_MLOR_SHELF_COL, value="Manufacutre")  # source's own typo
    mlor.cell(row=1, column=_MLOR_MLOR_COL, value="Current MLOR")
    for r2, code in enumerate(codes):
        row = _MLOR_FIRST_DATA_ROW + r2
        srow = sku_by_code[code]
        mlor.cell(row=row, column=_MLOR_CODE_COL, value=int(code))
        mlor.cell(row=row, column=_MLOR_SHELF_COL, value=int(srow["shelf_life_days"]))
        mlor.cell(row=row, column=_MLOR_MLOR_COL, value=int(srow["mlor_days"]))

    wb.save(path)
    return path


if __name__ == "__main__":
    import sys
    out = sys.argv[1] if len(sys.argv) > 1 else "synth_pet_workbook.xlsx"
    write_synth_workbook(out)
    print("wrote", out)
