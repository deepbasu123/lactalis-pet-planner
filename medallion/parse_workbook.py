"""medallion/parse_workbook.py

Parses a real "PET Traffic Lights" workbook (the Lactalis SNP/APO export plus
the manually-maintained production plan) into the six Silver-shaped tables the
medallion pipeline conforms: sku, week, parameter, demand, plan_line,
opening_stock.

This is the pipeline-side parser. It is deliberately SELF-CONTAINED (no import
of backend/*), because the FastAPI app is being reduced to a thin serving layer
and its excel_source.py may go away — the medallion pipeline must not depend on
it. The logic is lifted verbatim from the proven backend/excel_source.py; the
only change is that the parameter constants are inlined here rather than
imported from backend.data_gen.

Workbook shape (positions, not header text — see backend/excel_source.py for the
full provenance notes):

  "PET " sheet (trailing space is real)
    Row 13, col I (9) onward -> week_commencing dates (one column per week)
    Row 12, col I (9) onward -> "MAINT" marker for maintenance weeks
    Rows 15-25, col A/B/C     -> Status / Sku # / Sku Description (11 active SKUs)
    Rows 15-25, col I onward  -> planned production qty (EA) per week
  "SNP Dump" sheet
    Col B (2) SKU code, C (3) measure, E (5) location ("Total"),
    G (7) Initial (opening stock), H (8) onward one column per week.
  "MLOR" sheet
    Col A (1)/C (3)/E (5) -> product code / shelf life (days) / current MLOR (days)
"""
from __future__ import annotations

import datetime
import logging
import re
from typing import Any

import openpyxl
import pandas as pd

logger = logging.getLogger(__name__)

# ---------------------------------------------------------------------------
# Business-rule parameter constants (RULE-005: configurable, not per-refresh
# data). Inlined here so the parser is self-contained; kept byte-identical to
# backend.data_gen._PARAM_ROWS.
# ---------------------------------------------------------------------------
_PARAM_ROWS: list[tuple] = [
    ("cap_400ml_1_2_sku",      650_000.0, "400ml line capacity for 1-2 SKUs per week"),
    ("cap_400ml_3_sku",        600_000.0, "400ml line capacity for 3 SKUs per week"),
    ("cap_500ml_changeover",   650_000.0, "500ml line capacity for the first week after a changeover"),
    ("cap_500ml_steady_1_2",   700_000.0, "500ml line steady-state capacity for 1-2 SKUs"),
    ("cap_500ml_3sku_penalty",  50_000.0, "Capacity penalty for running a third 500ml SKU per week"),
    ("qa_hold_weeks",               2.0,  "QA hold period in weeks before production is sellable"),
    ("max_cover_weeks",            16.0,  "Cover weeks ceiling; stock beyond this fires the black band"),
    ("target_cover_weeks",          3.0,  "Planner target stock cover weeks"),
    ("reaction_window",             3.0,  "Weeks within which a stockout fires dark-red"),
    ("demand_horizon",              4.0,  "Forward demand weeks used in cover calculation"),
]

_PET_SHEET = "PET "
_SNP_SHEET = "SNP Dump"
_MLOR_SHEET = "MLOR"

_PET_SKU_ROWS = range(15, 26)
_PET_STATUS_COL = 1
_PET_CODE_COL = 2
_PET_DESC_COL = 3
_PET_WEEK_DATE_ROW = 13
_PET_MAINT_ROW = 12
_PET_WEEK_FIRST_COL = 9
_PET_MAINT_MARKER = "MAINT"
_PET_ACTIVE_STATUS = "Active"

_SNP_CODE_COL = 2
_SNP_MEASURE_COL = 3
_SNP_LOCATION_COL = 5
_SNP_INITIAL_COL = 7
_SNP_WEEK_FIRST_COL = 8
_SNP_LOCATION_TOTAL = "Total"

_MLOR_CODE_COL = 1
_MLOR_SHELF_COL = 3
_MLOR_MLOR_COL = 5
_MLOR_FIRST_DATA_ROW = 4

_HORIZON_WEEKS = 52
_LOCK_WEEKS = 3

# Shelf-life/MLOR imputation for SKUs absent from a real MLOR sheet: each maps
# to the closest real analog in the same file (same pack + family). Never used
# when the sheet already covers every SKU (as the synthetic workbook does).
_MLOR_FALLBACK: dict[str, str] = {
    "228500": "60444",
    "228510": "60444",
    "230540": "60444",
    "230550": "60444",
    "230150": "70526",
}

_DEMAND_MEASURES: dict[str, str] = {
    "forecast": "Forecast",
    "sales_order": "Sales Order",
    "distr_demand_planned": "DistrDemand (Planned)",
    "distr_demand_tlb": "DistrDemand (TLB Confirmed)",
}
_SOH_MEASURE = "Stock on Hand (Projected)"

_PACK_SIZE_RE = re.compile(r"(400|500)\s*ml", re.IGNORECASE)


def _clean_description(desc: Any) -> str:
    if desc is None:
        return ""
    text = str(desc).replace("\xa0", " ")
    text = re.sub(r"\s+", " ", text).strip()
    return text.rstrip("*").strip()


def infer_pack_size_ml(description: str) -> int:
    m = _PACK_SIZE_RE.search(description)
    if not m:
        raise ValueError(f"Could not infer pack size (400/500 ml) from description: {description!r}")
    return int(m.group(1))


def iso_week_key(d: datetime.date) -> str:
    y, w, _ = d.isocalendar()
    return f"{y}-W{w:02d}"


def _as_date(v: Any) -> datetime.date:
    if isinstance(v, datetime.datetime):
        return v.date()
    if isinstance(v, datetime.date):
        return v
    raise TypeError(f"Expected a date/datetime cell, got {v!r}")


def _num(v: Any) -> float:
    if v is None or v == "":
        return 0.0
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        return float(v)
    raise ValueError(f"Expected a blank cell or a number, got {v!r}")


def _parse_snp_date(v: Any) -> datetime.date:
    if isinstance(v, (datetime.datetime, datetime.date)):
        return _as_date(v)
    return datetime.datetime.strptime(str(v), "%d.%m.%Y").date()


def _read_sku_rows(pet_ws) -> list[tuple[int, str, str, int]]:
    rows: list[tuple[int, str, str, int]] = []
    priority = 0
    for r in _PET_SKU_ROWS:
        code = pet_ws.cell(row=r, column=_PET_CODE_COL).value
        if code is None:
            continue
        status = pet_ws.cell(row=r, column=_PET_STATUS_COL).value
        if str(status).strip().lower() != _PET_ACTIVE_STATUS.lower():
            logger.info("Skipping SKU %s: status is %r, not Active", code, status)
            continue
        priority += 1
        desc = pet_ws.cell(row=r, column=_PET_DESC_COL).value
        rows.append((priority, str(int(code)), _clean_description(desc), r))
    return rows


def _validate_week_alignment(snp_ws, week_dates: list[datetime.date]) -> None:
    for i, expected in enumerate(week_dates):
        raw = snp_ws.cell(row=1, column=_SNP_WEEK_FIRST_COL + i).value
        try:
            actual = _parse_snp_date(raw)
        except (TypeError, ValueError) as exc:
            raise ValueError(
                f"SNP Dump week column offset {i} (expected date "
                f"{expected.isoformat()}) has an unparseable header {raw!r}"
            ) from exc
        if actual != expected:
            raise ValueError(
                f"Week alignment mismatch at offset {i}: 'PET ' sheet says "
                f"{expected.isoformat()}, 'SNP Dump' says {actual.isoformat()}. "
                "Refusing to guess -- demand cannot be safely aligned to the week horizon."
            )


def _read_mlor(mlor_ws) -> dict[str, tuple[int, int]]:
    result: dict[str, tuple[int, int]] = {}
    for r in range(_MLOR_FIRST_DATA_ROW, mlor_ws.max_row + 1):
        code = mlor_ws.cell(row=r, column=_MLOR_CODE_COL).value
        if code is None:
            continue
        shelf = mlor_ws.cell(row=r, column=_MLOR_SHELF_COL).value
        mlor_days = mlor_ws.cell(row=r, column=_MLOR_MLOR_COL).value
        if shelf is None or mlor_days is None:
            continue
        result[str(int(code))] = (int(shelf), int(mlor_days))
    return result


def _read_week_dates(pet_ws, horizon_weeks: int) -> list[datetime.date]:
    dates = []
    for i in range(horizon_weeks):
        col = _PET_WEEK_FIRST_COL + i
        v = pet_ws.cell(row=_PET_WEEK_DATE_ROW, column=col).value
        dates.append(_as_date(v))
    return dates


def _read_maintenance_offsets(pet_ws, horizon_weeks: int) -> set[int]:
    offsets: set[int] = set()
    for i in range(horizon_weeks):
        col = _PET_WEEK_FIRST_COL + i
        if pet_ws.cell(row=_PET_MAINT_ROW, column=col).value == _PET_MAINT_MARKER:
            offsets.add(i)
    return offsets


def _index_snp_rows(snp_ws, codes: list[str]) -> dict[tuple[str, Any], int]:
    code_ints = {int(c) for c in codes}
    index: dict[tuple[str, Any], int] = {}
    for r in range(1, snp_ws.max_row + 1):
        code_val = snp_ws.cell(row=r, column=_SNP_CODE_COL).value
        if code_val not in code_ints:
            continue
        if snp_ws.cell(row=r, column=_SNP_LOCATION_COL).value != _SNP_LOCATION_TOTAL:
            continue
        measure = snp_ws.cell(row=r, column=_SNP_MEASURE_COL).value
        key = (str(code_val), measure)
        if key in index:
            raise ValueError(
                f"SNP Dump has duplicate {_SNP_LOCATION_TOTAL!r} rows for SKU "
                f"{code_val} / measure {measure!r} (rows {index[key]} and {r})."
            )
        index[key] = r
    return index


def _snp_week_values(snp_ws, row: int, horizon_weeks: int) -> list[float]:
    return [
        _num(snp_ws.cell(row=row, column=_SNP_WEEK_FIRST_COL + i).value)
        for i in range(horizon_weeks)
    ]


def _build_sku_df(sku_rows, mlor_by_code) -> pd.DataFrame:
    records = []
    for priority, code, desc, _row in sku_rows:
        pack_ml = infer_pack_size_ml(desc)
        if code in mlor_by_code:
            shelf, mlor_days = mlor_by_code[code]
        else:
            fallback_code = _MLOR_FALLBACK[code]
            shelf, mlor_days = mlor_by_code[fallback_code]
            logger.info("SKU %s has no MLOR entry -- imputing from SKU %s", code, fallback_code)
        records.append({
            "sku_code": code,
            "description": desc,
            "pack_size_ml": pack_ml,
            "priority": priority,
            "status": "Active",
            "shelf_life_days": shelf,
            "mlor_days": mlor_days,
            "max_cover_weeks": round((shelf - mlor_days) / 7.0, 1),
        })
    return pd.DataFrame(records)


def _build_week_df(week_dates, maint_offsets) -> pd.DataFrame:
    records = []
    for i, d in enumerate(week_dates):
        hi = i + 1
        records.append({
            "week_key": iso_week_key(d),
            "horizon_index": hi,
            "week_commencing": d,
            "maintenance_type": "Full" if i in maint_offsets else "None",
            "is_locked": hi <= _LOCK_WEEKS,
            "note": "",
        })
    return pd.DataFrame(records)


def _build_demand_df(snp_ws, codes, week_keys) -> pd.DataFrame:
    index = _index_snp_rows(snp_ws, codes)
    horizon_weeks = len(week_keys)
    records = []
    for code in codes:
        series: dict[str, list[float]] = {}
        for field, measure in _DEMAND_MEASURES.items():
            row = index.get((code, measure))
            if row is None:
                raise KeyError(f"SNP Dump has no {measure!r} / Total row for SKU {code}")
            series[field] = _snp_week_values(snp_ws, row, horizon_weeks)
        for i, wk in enumerate(week_keys):
            records.append({
                "sku_code": code,
                "week_key": wk,
                "forecast": series["forecast"][i],
                "sales_order": series["sales_order"][i],
                "distr_demand_planned": series["distr_demand_planned"][i],
                "distr_demand_tlb": series["distr_demand_tlb"][i],
            })
    return pd.DataFrame(records)


def _build_opening_df(snp_ws, codes) -> pd.DataFrame:
    index = _index_snp_rows(snp_ws, codes)
    records = []
    for code in codes:
        row = index.get((code, _SOH_MEASURE))
        if row is None:
            raise KeyError(f"SNP Dump has no {_SOH_MEASURE!r} / Total row for SKU {code}")
        initial = _num(snp_ws.cell(row=row, column=_SNP_INITIAL_COL).value)
        records.append({"sku_code": code, "opening_ea": initial})
    return pd.DataFrame(records)


def _build_plan_df(pet_ws, sku_rows, week_keys) -> pd.DataFrame:
    records = []
    for _priority, code, _desc, r in sku_rows:
        for i, wk in enumerate(week_keys):
            qty = _num(pet_ws.cell(row=r, column=_PET_WEEK_FIRST_COL + i).value)
            records.append({"sku_code": code, "week_key": wk, "planned_qty": qty, "orig_qty": qty})
    return pd.DataFrame(records)


def load_dataset(path: str, horizon_weeks: int = _HORIZON_WEEKS) -> dict[str, pd.DataFrame]:
    """Parse a "PET Traffic Lights" workbook into the six Silver-shaped tables."""
    wb = openpyxl.load_workbook(path, data_only=True)
    for sheet in (_PET_SHEET, _SNP_SHEET, _MLOR_SHEET):
        if sheet not in wb.sheetnames:
            raise KeyError(f"Workbook is missing the expected {sheet!r} sheet")

    pet = wb[_PET_SHEET]
    snp = wb[_SNP_SHEET]
    mlor = wb[_MLOR_SHEET]

    sku_rows = _read_sku_rows(pet)
    if not sku_rows:
        raise ValueError(f"No active SKUs found in {_PET_SHEET!r} rows {list(_PET_SKU_ROWS)}")

    mlor_by_code = _read_mlor(mlor)
    week_dates = _read_week_dates(pet, horizon_weeks)
    week_keys = [iso_week_key(d) for d in week_dates]
    maint_offsets = _read_maintenance_offsets(pet, horizon_weeks)
    codes = [code for _, code, _, _ in sku_rows]

    _validate_week_alignment(snp, week_dates)

    logger.info("Parsed %d SKUs and %d weeks (%s to %s) from %s",
                len(sku_rows), horizon_weeks, week_keys[0], week_keys[-1], path)

    return {
        "sku": _build_sku_df(sku_rows, mlor_by_code),
        "week": _build_week_df(week_dates, maint_offsets),
        "parameter": pd.DataFrame(
            [{"name": n, "value": v, "description": d} for n, v, d in _PARAM_ROWS]
        ),
        "demand": _build_demand_df(snp, codes, week_keys),
        "plan_line": _build_plan_df(pet, sku_rows, week_keys),
        "opening_stock": _build_opening_df(snp, codes),
    }
